"""Offline lead QA. No originals overwritten; no automatic employment claims."""
import io
import re
import unicodedata
from datetime import date
from urllib.parse import urlsplit

import pandas as pd
import phonenumbers
import pycountry
import tldextract
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EXTRACT = tldextract.TLDExtract(suffix_list_urls=(), include_psl_private_domains=True)
ALIASES = {
 'company': ['companyname', 'company name', 'company', 'organization', 'organisation'],
 'email': ['email', 'email address', 'work email'],
 'website': ['website', 'company website', 'companyurl', 'domain', 'url'],
 'country': ['lead country', 'country', 'contact country'],
 'region': ['c state', 'state', 'region', 'province'],
 'industry': ['c industry', 'companyindustry', 'industry'],
 'department': ['departments', 'department', 'job department', 'function'],
 'job_level': ['job level', 'seniority', 'level'],
 'job_role': ['job role c', 'job role', 'role'],
 'companysize': ['number of employees', 'employees count', 'company size', 'companysize', 'employees'],
 'phone': ['telephone', 'phone', 'phone 1', 'mobile', 'direct', 'switch'],
 'jobtitle': ['jobtitle', 'job title', 'title'],
}
PICK_FIELDS = ['country', 'region', 'industry', 'department', 'job_level', 'job_role', 'companysize']
REL_COLS = ['company', 'domain', 'relationship', 'evidence_url', 'checked_date']
REL_TYPES = {'official', 'trading_name', 'rebrand', 'parent', 'subsidiary', 'acquisition', 'unrelated'}
COLORS = {'green': 'C6EFCE', 'amber': 'F8CBAD', 'yellow': 'FFF2CC', 'red': 'FFC7CE', 'grey': 'E7E6E6'}

def text(v):
    return '' if v is None or pd.isna(v) else str(v).strip()

def key(v):
    s = unicodedata.normalize('NFKC', text(v)).casefold().replace('&', ' and ')
    return re.sub(r'\s+', ' ', re.sub(r'[^\w]+', ' ', s.replace('_', ' '))).strip()

def detect(df):
    """Exact normalized aliases only; never guess an unrelated header."""
    out = {}
    for field, aliases in ALIASES.items():
        out[field] = None
        for alias in aliases:
            candidates = [c for c in df.columns if key(c) == key(alias)]
            if len(candidates) == 1:
                out[field] = candidates[0]
                break
    return out

def domain(v, email=False):
    s = text(v).lower()
    if email:
        if s.count('@') != 1 or not s.split('@')[0] or re.search(r'\s', s):
            return ''
        s = s.split('@')[1]
    try:
        host = urlsplit(s if '://' in s else 'https://' + s).hostname or ''
        host = host.rstrip('.').encode('idna').decode('ascii')
    except (ValueError, UnicodeError):
        return ''
    if not re.fullmatch(r'[a-z0-9.-]+', host) or '..' in host:
        return ''
    ext = EXTRACT(host)
    return ext.top_domain_under_public_suffix if ext.suffix and ext.domain else ''

def read_frame(data, sheet):
    df = pd.read_excel(io.BytesIO(data), sheet_name=sheet, dtype=str, keep_default_na=False)
    if any(str(c).startswith('Unnamed:') for c in df.columns):
        df = df.loc[:, ~df.columns.astype(str).str.startswith('Unnamed:')]
    return df

def relationships(data=None, filename=''):
    if not data:
        return {}
    df = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False) if filename.lower().endswith('.csv') else read_frame(data, 0)
    if not set(REL_COLS).issubset(df.columns):
        raise ValueError('Relationship file needs columns: ' + ', '.join(REL_COLS))
    records = {}
    for _, row in df.iterrows():
        if not any(text(row[c]) for c in REL_COLS):
            continue
        c, d, rel = key(row['company']), domain(row['domain']), text(row['relationship']).lower()
        url = text(row['evidence_url'])
        try:
            checked = date.fromisoformat(text(row['checked_date']))
        except ValueError as exc:
            raise ValueError('Relationship checked_date must be YYYY-MM-DD.') from exc
        if not c or not d or rel not in REL_TYPES or not url.startswith(('https://', 'http://')) or checked > date.today():
            raise ValueError('Invalid relationship row. Check company, domain, type, evidence URL and date.')
        if (c, d) in records:
            raise ValueError(f'Duplicate company/domain relationship: {c}, {d}')
        records[c, d] = (rel, url, checked)
    return records

def company_check(company, dom, rels, max_age=180):
    if not text(company) or not dom:
        return ('Unable to check', 'grey', 'Missing company or valid domain', '', '')
    saved = rels.get((key(company), dom))
    if saved:
        rel, url, checked = saved
        if (date.today() - checked).days > max_age:
            return ('Evidence needs refresh', 'yellow', rel, url, checked.isoformat())
        if rel in {'official', 'trading_name'}:
            return ('Recognised company domain', 'green', rel, url, checked.isoformat())
        if rel == 'unrelated':
            return ('Company/email conflict', 'red', 'Reviewed as unrelated', url, checked.isoformat())
        return ('Connected organisation/domain', 'amber', rel + '; confirm email applicability', url, checked.isoformat())
    toks = [t for t in key(company).split() if t not in {'and','the','llp','llc','ltd','limited','inc','pllc','corp','company','group'}]
    base = EXTRACT(dom).domain
    compact = ''.join(toks)
    acronym = ''.join(t[0] for t in toks)
    similar = bool(compact) and (compact == base or (len(acronym) >= 2 and acronym == base) or fuzz.ratio(compact, base) >= 80)
    if similar:
        return ('Possible match — unverified', 'yellow', 'Name similarity only; relationship not verified', '', '')
    return ('Potential mismatch — research needed', 'yellow', 'Name differs; no verified relationship recorded', '', '')

def country_iso(value):
    alias = {'uk': 'GB', 'usa': 'US', 'united states of america': 'US', 'south korea': 'KR'}
    if text(value).lower() in alias:
        return alias[text(value).lower()]
    try:
        return pycountry.countries.lookup(text(value)).alpha_2
    except LookupError:
        return None

def phone_check(value, country):
    s = text(value)
    if not s:
        return 'Not checked: missing phone'
    region = country_iso(country)
    s = re.sub(r'\.0+$', '', s)
    if s.startswith('00'):
        s = '+' + s[2:]
    if not s.startswith('+') and not region:
        return 'Not checked: unknown country'
    try:
        num = phonenumbers.parse(s, None if s.startswith('+') else region)
    except phonenumbers.NumberParseException:
        return 'Review: cannot parse phone'
    if not phonenumbers.is_valid_number(num):
        return 'Review: invalid phone format'
    if not region:
        return 'Not checked: unknown country'
    return 'Valid format / country match' if phonenumbers.is_valid_number_for_region(num, region) else 'Review: phone country differs'

def placeholder(v):
    return bool(re.match(r'^(?:leave blank|picklist|integer|text|mandatory|do not map|all accepted.*|no proof:.*|\+?x+|dd/mm/yyyy.*)$', text(v), re.I))

def process(master, sheet, mapping, pick=None, pick_sheet=0, pick_mapping=None, rels=None, skip_template=True, colours=True):
    df = read_frame(master, sheet)
    rels = rels or {}
    pdf = read_frame(pick, pick_sheet) if pick else pd.DataFrame()
    pm = pick_mapping or detect(pdf)
    allowed = {f: {key(v) for v in pdf[pm[f]] if text(v) and not placeholder(v)} if pm.get(f) in pdf.columns else set() for f in PICK_FIELDS}
    wb = load_workbook(io.BytesIO(master))
    ws = wb[sheet]
    headers = {text(c.value): c.column for c in ws[1] if c.value is not None}
    rows = []
    for i, row in df.iterrows():
        values = {f: text(row[col]) if col in df.columns else '' for f, col in mapping.items()}
        result, fills = {}, {}
        def put(name, value, color='grey'):
            result['QA_' + name] = value
            fills['QA_' + name] = color
        is_blank = not any(text(v) for c,v in row.items() if not str(c).startswith('QA_'))
        is_template = skip_template and sum(placeholder(v) for c,v in row.items() if not str(c).startswith('QA_')) >= 3
        if is_blank or is_template:
            put('Overall_Status', 'SKIPPED: blank/template row')
        else:
            company = values.get('company', '')
            ed, wd = domain(values.get('email'), True), domain(values.get('website'))
            put('Email_Domain', ed)
            put('Website_Domain', wd)
            for label, dom in [('Company_Email', ed), ('Company_Website', wd)]:
                status, color, why, url, checked = company_check(company, dom, rels)
                put(label + '_Status', status, color)
                put(label + '_Reason', why)
                put(label + '_Evidence_URL', url)
                put(label + '_Checked_Date', checked)
            if not ed or not wd:
                put('Email_Website_Status', 'Unable to check', 'grey')
            elif ed == wd:
                put('Email_Website_Status', 'Same domain (not employment verification)', 'green')
            elif all((key(company), d) in rels and (date.today()-rels[key(company),d][2]).days <= 180 and rels[key(company),d][0] != 'unrelated' for d in [ed,wd]):
                put('Email_Website_Status', 'Both have recorded company relationships — review', 'amber')
            else:
                put('Email_Website_Status', 'Different domains — review', 'yellow')
            put('Phone_Check', phone_check(values.get('phone'), values.get('country')))
            for f in PICK_FIELDS:
                if not mapping.get(f):
                    status = 'Not checked: master column missing'
                elif not allowed[f]:
                    status = 'Not checked: no picklist values'
                else:
                    status = 'Match' if key(values.get(f)) in allowed[f] else 'Review: not in picklist'
                put('Picklist_' + f, status, 'green' if status == 'Match' else 'yellow' if status.startswith('Review') else 'grey')
            missing = [f for f in ['company','email','country'] if not values.get(f)]
            put('Missing_Required_Fields', ', '.join(missing), 'yellow' if missing else 'green')
            problems = [name.removeprefix('QA_') for name, val in result.items() if (name.endswith('_Status') and fills[name] != 'green') or (name.startswith('QA_Picklist_') and val != 'Match') or (name == 'QA_Phone_Check' and val != 'Valid format / country match')]
            if missing:
                problems.append('Missing_Required_Fields')
            put('Overall_Status', 'REVIEW' if problems else 'CHECKS PASSED — employment/email not verified', 'yellow' if problems else 'green')
            put('Review_Reasons', '; '.join(problems))
            put('Recommended_Action', 'Verify company/domain relationships and current employment; do not auto-correct or discard.' if problems else 'Relationship checks passed; separate employment and mailbox verification still required.')
        # Clear previous QA results on reruns, including skipped rows.
        for h,c in headers.items():
            if h.startswith('QA_'):
                ws.cell(i+2,c).value = None
                ws.cell(i+2,c).fill = PatternFill(fill_type=None)
        for name,value in result.items():
            if name not in headers:
                headers[name] = ws.max_column + 1
                ws.cell(1, headers[name], name)
                ws.column_dimensions[ws.cell(1,headers[name]).column_letter].width = 28
            cell = ws.cell(i+2, headers[name])
            cell.value = value
            if isinstance(value,str):
                cell.data_type = 's'  # prevent formula injection in evidence/notes
            if colours:
                cell.fill = PatternFill('solid', fgColor=COLORS[fills[name]])
        rows.append(result)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), pd.DataFrame(rows)
